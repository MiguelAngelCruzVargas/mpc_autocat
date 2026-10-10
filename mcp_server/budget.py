"""Catálogo de conceptos para el presupuesto, SIN precios.

La herramienta ya sabe qué se dibujó y cuánto mide (calculate_quantities).
Lo que le ahorra trabajo a quien presupuesta es tener los CONCEPTOS
redactados, con su unidad y su cantidad, en una hoja donde solo escribe el
precio unitario: el importe, los subtotales por partida, el IVA y el total
ya están formulados.

Dos reglas de oficio que el cuadro de volúmenes no sigue y el catálogo sí:

- Un concepto es un TRABAJO con su cantidad NETA. La merma va adentro del
  precio unitario (es costo del contratista, no obra ejecutada): el
  catálogo descuenta la merma que traen los items.
- Los insumos no son conceptos. Piezas de block, mortero de junteo,
  cemento/arena/grava de la mezcla van INCLUIDOS en el P.U. del muro o del
  concreto: el muro se presupuesta en m², no en piezas, y la pintura se mide
  en m² de superficie ("a dos manos" va en la descripción, no multiplica la
  cantidad).

Nada se inventa: la cantidad sale de lo medido. Un concepto que hace falta
y que el dibujo no mide (la cimentación sin estudio, por ejemplo) entra con
la cantidad en blanco y marcado POR CUANTIFICAR.
"""
from __future__ import annotations

import csv
import io
import os
from typing import Any, Optional

INCLUYE = "Incluye: materiales, mano de obra, herramienta y equipo."

# Tipos de calculate_quantities que son INSUMOS, no conceptos.
INSUMOS = {"mortar_volume", "concrete_mix"}


ORDEN_OBRA = ["preliminar", "demolic", "excavac", "terracer", "cimentac",
              "estructura", "albanil", "instalac", "herrer",
              "carpinter", "cancel", "acabado", "limpieza"]


def _normal(texto: str) -> str:
    t = str(texto).strip().lower()
    for a, b in (("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u"), ("ñ", "n")):
        t = t.replace(a, b)
    return t


def _neto(valor: float, merma_pct: float) -> float:
    return float(valor) / (1.0 + float(merma_pct or 0.0) / 100.0)


def _sin_punto(texto: str) -> str:
    return str(texto).strip().rstrip(".")


def concepto_de_item(item: dict[str, Any]) -> Optional[dict[str, Any]]:
    """El concepto de presupuesto que corresponde a un item medido, o None
    si el item es un insumo (va adentro del P.U. de otro concepto)."""
    tipo = item["type"]
    if tipo in INSUMOS:
        return None
    label = _sin_punto(item["label"])
    merma = item.get("wastePct", 0.0)

    if tipo == "concrete_volume":
        unidad, cantidad = "m3", _neto(item["volume"], merma)
        texto = (f"{label}. Incluye: suministro de materiales, elaboración, "
                 "colado, vibrado y curado; mano de obra, herramienta y equipo.")
    elif tipo == "steel_weight":
        detalle = item.get("detail") or {}
        # Solo la MALLA electrosoldada se presupuesta por m²; una parrilla
        # de varilla cuantificada por kg/m² sigue siendo acero de refuerzo
        # (kg). 'concept_unit' lo fuerza si hace falta.
        malla = (set(detalle) == {"mesh"} and "malla" in label.lower())
        if item.get("concept_unit"):
            malla = str(item["concept_unit"]).lower() in ("m2", "m²")
        if malla and "mesh" in detalle:
            unidad, cantidad = "m2", float(detalle["mesh"]["area"])
            texto = (f"Suministro y colocación de {label[0].lower() + label[1:]}. "
                     "Incluye: traslapes, amarres, calzas, mano de obra, "
                     "herramienta y equipo.")
        else:
            unidad, cantidad = "kg", _neto(item["weight"], merma)
            texto = (f"Suministro, habilitado y armado de "
                     f"{label[0].lower() + label[1:]}. Incluye: cortes, "
                     "dobleces, traslapes, alambre recocido, desperdicios, "
                     "mano de obra, herramienta y equipo.")
    elif tipo == "brick_count":
        unidad, cantidad = "m2", float(item["wallArea"])
        texto = (f"{label}, asentado con mortero, a plomo y nivel. Incluye: "
                 "materiales, acarreos, desperdicios, mano de obra, "
                 "herramienta y equipo.")
    elif tipo == "area_finish":
        # 'area' ya trae las caras; las manos van en la descripción.
        unidad, cantidad = "m2", float(item["area"])
        texto = (f"{label}. Incluye: preparación de la superficie, "
                 "materiales, mano de obra, herramienta y equipo.")
    elif tipo == "formwork":
        unidad, cantidad = "m2", _neto(item["areaTotal"], merma)
        texto = (f"Cimbra y descimbra de {label[0].lower() + label[1:]}. "
                 "Incluye: materiales, habilitado, mano de obra, herramienta "
                 "y equipo.")
    elif tipo == "earthwork":
        unidad, cantidad = "m3", float(item["volume"])
        texto = (f"{label}. Incluye: mano de obra, herramienta, equipo y "
                 "acarreo dentro de la obra.")
    else:
        unidad, cantidad = item.get("unit", ""), None
        texto = f"{label}. {INCLUYE}"

    return {"group": item.get("group") or "General",
            "concept": item.get("concept") or texto,
            "unit": unidad, "quantity": cantidad,
            "source": item.get("label", "")}


def armar(result: dict[str, Any],
          extra_concepts: Optional[list[dict[str, Any]]] = None
          ) -> dict[str, Any]:
    """Partidas con sus conceptos (y claves), a partir del resultado de
    calculate_quantities y de los conceptos extra que no salen del dibujo."""
    conceptos, insumos = [], []
    for item in result.get("items", []):
        c = concepto_de_item(item)
        if c is None:
            insumos.append(item["label"])
        else:
            conceptos.append(c)
    for i, e in enumerate(extra_concepts or []):
        if not e.get("concept") or not e.get("unit"):
            raise ValueError(f"El concepto extra #{i + 1} necesita 'concept' "
                             "(la descripción) y 'unit'.")
        q = e.get("quantity")
        conceptos.append({"group": e.get("group") or "General",
                          "concept": _sin_punto(e["concept"]) + ".",
                          "unit": e["unit"],
                          "quantity": None if q is None else float(q),
                          "source": e.get("source", "")})

    # Las partidas van en el orden en que se ejecuta la obra; las que no
    # están en la lista quedan después, en el orden en que aparecieron.
    def _orden(c: dict[str, Any]) -> int:
        nombre = _normal(c["group"])
        return next((i for i, g in enumerate(ORDEN_OBRA) if nombre.startswith(g)),
                    len(ORDEN_OBRA))
    conceptos = sorted(conceptos, key=_orden)   # estable: respeta el resto

    partidas: list[dict[str, Any]] = []
    for c in conceptos:
        p = next((p for p in partidas if p["name"] == c["group"]), None)
        if p is None:
            p = {"name": c["group"], "key": f"{len(partidas) + 1:02d}",
                 "concepts": []}
            partidas.append(p)
        c["key"] = f"{p['key']}.{len(p['concepts']) + 1:02d}"
        p["concepts"].append(c)
    return {"partidas": partidas, "insumosIncluidos": insumos,
            "porCuantificar": [c["key"] for p in partidas
                               for c in p["concepts"] if c["quantity"] is None]}


def export_concept_catalog(result: dict[str, Any], path: str,
                           extra_concepts: Optional[list[dict[str, Any]]] = None,
                           project: str = "", vat_pct: float = 16.0
                           ) -> dict[str, Any]:
    """Escribe el catálogo en .xlsx (formulado) o .csv. Ver el módulo."""
    if not path:
        raise ValueError("Falta 'path' (termina en .xlsx o .csv).")
    cat = armar(result, extra_concepts)
    if not cat["partidas"]:
        raise ValueError("No hay conceptos: el resultado no trae items que "
                         "sean trabajos (solo insumos o nada).")
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    if path.lower().endswith(".csv"):
        _csv(cat, path)
    elif path.lower().endswith(".xlsx"):
        _xlsx(cat, path, project, vat_pct)
    else:
        raise ValueError("El catálogo se escribe en .xlsx (con fórmulas) o .csv.")

    n = sum(len(p["concepts"]) for p in cat["partidas"])
    salida: dict[str, Any] = {
        "path": os.path.abspath(path), "partidas": len(cat["partidas"]),
        "concepts": n,
        "catalog": [{"key": c["key"], "concept": c["concept"],
                     "unit": c["unit"],
                     "quantity": None if c["quantity"] is None
                     else round(c["quantity"], 2)}
                    for p in cat["partidas"] for c in p["concepts"]],
    }
    if cat["insumosIncluidos"]:
        salida["includedSupplies"] = cat["insumosIncluidos"]
        salida["note"] = ("No van como concepto (son insumos, entran en el "
                          "P.U.): " + ", ".join(cat["insumosIncluidos"]) + ".")
    if cat["porCuantificar"]:
        salida["warning"] = ("Conceptos POR CUANTIFICAR (sin cantidad medida): "
                             + ", ".join(cat["porCuantificar"]) + ".")
    return salida


def _csv(cat: dict[str, Any], path: str) -> None:
    with io.open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["CLAVE", "CONCEPTO", "UNIDAD", "CANTIDAD", "P.U.", "IMPORTE"])
        for p in cat["partidas"]:
            w.writerow([p["key"], p["name"].upper(), "", "", "", ""])
            for c in p["concepts"]:
                q = "" if c["quantity"] is None else f"{c['quantity']:.2f}"
                w.writerow([c["key"], c["concept"], c["unit"],
                            q or "POR CUANTIFICAR", "", ""])


def _xlsx(cat: dict[str, Any], path: str, project: str, vat_pct: float) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.properties import PageSetupProperties

    wb = Workbook()
    ws = wb.active
    ws.title = "Catálogo"
    delgado = Side(style="thin", color="808080")
    borde = Border(left=delgado, right=delgado, top=delgado, bottom=delgado)
    gris = PatternFill("solid", fgColor="D9D9D9")
    claro = PatternFill("solid", fgColor="F2F2F2")
    entrada = PatternFill("solid", fgColor="FFF2CC")   # donde se escribe
    dinero = '"$"#,##0.00'

    for col, ancho in zip("ABCDEF", (9, 70, 8, 12, 14, 16)):
        ws.column_dimensions[col].width = ancho
    ws["A1"] = "CATÁLOGO DE CONCEPTOS"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = project or ""
    ws["A2"].font = Font(bold=True)
    ws["A3"] = ("Cantidades netas medidas del plano. Llenar solo P.U. "
                "(celdas amarillas): el importe y los totales se calculan.")
    ws["A3"].font = Font(italic=True, color="595959")

    fila = 5
    for col, txt in zip("ABCDEF", ("CLAVE", "CONCEPTO", "UNIDAD", "CANTIDAD",
                                   "P.U.", "IMPORTE")):
        c = ws[f"{col}{fila}"]
        c.value, c.font, c.fill, c.border = txt, Font(bold=True), gris, borde
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = f"A{fila + 1}"

    subtotales = []
    for p in cat["partidas"]:
        fila += 1
        ws[f"A{fila}"], ws[f"B{fila}"] = p["key"], p["name"].upper()
        for col in "ABCDEF":
            ws[f"{col}{fila}"].fill, ws[f"{col}{fila}"].border = claro, borde
            ws[f"{col}{fila}"].font = Font(bold=True)
        primera = fila + 1
        for c in p["concepts"]:
            fila += 1
            ws[f"A{fila}"] = c["key"]
            ws[f"B{fila}"] = c["concept"]
            ws[f"C{fila}"] = c["unit"]
            if c["quantity"] is None:
                ws[f"D{fila}"] = "POR CUANTIFICAR"
                ws[f"D{fila}"].font = Font(color="C00000", bold=True)
            else:
                ws[f"D{fila}"] = round(c["quantity"], 2)
                ws[f"D{fila}"].number_format = "#,##0.00"
            ws[f"E{fila}"].fill = entrada
            ws[f"E{fila}"].number_format = dinero
            ws[f"F{fila}"] = f'=IF(ISNUMBER(D{fila}),ROUND(D{fila}*E{fila},2),0)'
            ws[f"F{fila}"].number_format = dinero
            ws[f"B{fila}"].alignment = Alignment(wrap_text=True, vertical="top")
            for col in "ACDEF":
                ws[f"{col}{fila}"].alignment = Alignment(vertical="top",
                                                         horizontal="center" if col in "AC" else "right")
            for col in "ABCDEF":
                ws[f"{col}{fila}"].border = borde
        fila += 1
        ws[f"E{fila}"] = f"SUBTOTAL {p['key']}"
        ws[f"E{fila}"].font = Font(bold=True)
        ws[f"F{fila}"] = f"=SUM(F{primera}:F{fila - 1})"
        ws[f"F{fila}"].number_format = dinero
        ws[f"F{fila}"].font = Font(bold=True)
        subtotales.append(f"F{fila}")

    fila += 2
    ws[f"E{fila}"], ws[f"F{fila}"] = "SUBTOTAL", "=" + "+".join(subtotales)
    base = fila
    if vat_pct:
        fila += 1
        ws[f"E{fila}"] = f"IVA {vat_pct:g}%"
        ws[f"F{fila}"] = f"=ROUND(F{base}*{vat_pct / 100.0:g},2)"
    fila += 1
    ws[f"E{fila}"] = "TOTAL"
    ws[f"F{fila}"] = f"=F{base}" + (f"+F{base + 1}" if vat_pct else "")
    for r in range(base, fila + 1):
        ws[f"E{r}"].font = ws[f"F{r}"].font = Font(bold=True)
        ws[f"F{r}"].number_format = dinero

    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = "5:5"
    wb.save(path)
