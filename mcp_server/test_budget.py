"""Tests del catálogo de conceptos (presupuesto sin precios), sin AutoCAD.

Lo que se fija acá son reglas de oficio que el cuadro de volúmenes no
sigue: cantidades NETAS (la merma va en el P.U.), insumos fuera del
catálogo, el muro en m² y no en piezas, la pintura por superficie y no por
manos, partidas en orden de obra, y nada inventado (POR CUANTIFICAR).

Uso:  python test_budget.py
"""
from __future__ import annotations

import os
import re
import sys
import tempfile

import budget

FAILED: list[str] = []


def check(name: str, ok: bool, detalle: str = "") -> None:
    if not ok:
        FAILED.append(f"{name}: {detalle}")
    print(("  ok  " if ok else " FALLA ") + name
          + (f" -- {detalle}" if not ok else ""))


RESULTADO = {"items": [
    {"label": "Muro de block 15x20x40 cm", "type": "brick_count",
     "areas": [145.573], "wallArea": 119.03, "wastePct": 3.0, "pieces": 1424,
     "group": "Albañilería"},
    {"label": "Mortero de junteo", "type": "mortar_volume", "wastePct": 3.0,
     "mortarVolume": 1.3, "group": "Albañilería"},
    {"label": "Pintura vinílica, 2 manos (dos caras)", "type": "area_finish",
     "areas": [145.573], "area": 291.146, "faces": 2, "coats": 2,
     "wastePct": 5.0, "billableArea": 611.41, "group": "Acabados"},
    {"label": "Losa f'c=200, h=0.12 m", "type": "concrete_volume",
     "volume": 7.4513, "wastePct": 5.0, "group": "Estructura"},
    {"label": "Acero de losa: varilla #3 @20 cm", "type": "steel_weight",
     "detail": {"mesh": {"area": 59.1375, "kgM2": 5.6}}, "weight": 347.73,
     "wastePct": 5.0, "group": "Estructura"},
    {"label": "Malla electrosoldada 6x6-10/10", "type": "steel_weight",
     "detail": {"mesh": {"area": 59.1375, "kgM2": 0.95}}, "weight": 58.99,
     "wastePct": 5.0, "group": "Estructura"},
    {"label": "Mezcla f'c=200", "type": "concrete_mix", "group": "Estructura"},
]}


def _catalogo(extra=None):
    with tempfile.TemporaryDirectory() as d:
        r = budget.export_concept_catalog(RESULTADO, os.path.join(d, "c.xlsx"),
                                          extra_concepts=extra)
        from openpyxl import load_workbook
        ws = load_workbook(os.path.join(d, "c.xlsx")).active
        filas = [r_ for r_ in ws.iter_rows(values_only=True)]
    return r, {c["concept"].split(".")[0]: c for c in r["catalog"]}, filas


def test_cantidades_netas_y_unidades_de_presupuesto() -> None:
    r, por, _ = _catalogo()
    muro = next(c for k, c in por.items() if "block" in k)
    check("el muro va en m2 (area neta), no en piezas",
          muro["unit"] == "m2" and muro["quantity"] == 119.03, str(muro))
    pintura = next(c for k, c in por.items() if "Pintura" in k)
    check("la pintura por superficie (2 caras), no multiplicada por manos",
          pintura["quantity"] == 291.15, str(pintura))
    losa = next(c for k, c in por.items() if "Losa" in k)
    check("concreto sin la merma (7.4513 / 1.05)",
          abs(losa["quantity"] - 7.4513 / 1.05) < 0.01, str(losa))
    acero = next(c for k, c in por.items() if "acero de losa" in k.lower())
    check("una parrilla de varilla sigue en kg", acero["unit"] == "kg", str(acero))
    malla = next(c for k, c in por.items() if "malla" in k.lower())
    check("la malla electrosoldada va en m2",
          malla["unit"] == "m2" and malla["quantity"] == 59.14, str(malla))


def test_insumos_fuera_del_catalogo() -> None:
    r, por, _ = _catalogo()
    check("mortero y mezcla no son conceptos",
          not any("Mortero" in k or "Mezcla" in k for k in por), str(list(por)))
    check("y se dice que van en el P.U.",
          "Mortero de junteo" in r.get("note", ""), str(r.get("note")))


def test_partidas_en_orden_de_obra_y_extras() -> None:
    extra = [{"group": "Preliminares", "concept": "Limpieza y trazo",
              "unit": "m2", "quantity": 105},
             {"group": "Cimentación", "concept": "Cimentación según estudio",
              "unit": "lote", "quantity": None}]
    r, por, filas = _catalogo(extra)
    claves = [c["key"] for c in r["catalog"]]
    check("preliminares primero (01.01)", r["catalog"][0]["concept"]
          .startswith("Limpieza") and claves[0] == "01.01", str(claves))
    nombres = [f[1] for f in filas if f[0] and len(str(f[0])) == 2]
    check("orden: preliminares, cimentación, estructura, albañilería, acabados",
          nombres == ["PRELIMINARES", "CIMENTACIÓN", "ESTRUCTURA",
                      "ALBAÑILERÍA", "ACABADOS"], str(nombres))
    sin = next(f for f in filas if f[1] and str(f[1]).startswith("Cimentación"))
    check("lo que no se mide sale POR CUANTIFICAR, sin número",
          sin[3] == "POR CUANTIFICAR" and "02.01" in r.get("warning", ""),
          str(sin))


def test_hoja_formulada_sin_precios() -> None:
    _, _, filas = _catalogo()
    conceptos = [f for f in filas if f[0] and re.fullmatch(r"\d\d\.\d\d", str(f[0]))]
    check("P.U. vacío en todos los conceptos",
          all(f[4] is None for f in conceptos), str([f[4] for f in conceptos]))
    check("importe formulado cantidad x P.U.",
          all(str(f[5]).startswith("=IF(ISNUMBER(D") for f in conceptos),
          str(conceptos[0][5]))
    textos = [str(f[4]) for f in filas if f[4]]
    check("subtotal, IVA y total",
          "SUBTOTAL" in textos and "TOTAL" in textos
          and any(t.startswith("IVA") for t in textos), str(textos[-4:]))


def test_errores_claros() -> None:
    for kw, frag in (({"path": "c.pdf"}, "xlsx"),
                     ({"path": ""}, "path"),
                     ({"path": "c.xlsx", "extra_concepts": [{"unit": "m2"}]},
                      "concept")):
        try:
            budget.export_concept_catalog(RESULTADO, **kw)
            check(f"error claro ({frag})", False, "no falló")
        except ValueError as exc:
            check(f"error claro ({frag})", frag in str(exc), str(exc))


def main() -> int:
    for t in (test_cantidades_netas_y_unidades_de_presupuesto,
              test_insumos_fuera_del_catalogo,
              test_partidas_en_orden_de_obra_y_extras,
              test_hoja_formulada_sin_precios,
              test_errores_claros):
        print(t.__name__)
        t()
    if FAILED:
        print(f"\n{len(FAILED)} FALLAS:")
        for f in FAILED:
            print(" - " + f)
        return 1
    print("\nOK: el catálogo trae conceptos de obra, netos y sin precios.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
